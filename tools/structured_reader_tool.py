"""파일 구조 분석 Tool — 제목, 표, 내용을 구분해서 읽기"""
import os
from langchain_core.tools import tool


def _parse_markdown(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    sections = []
    current_section = []

    for line in lines:
        stripped = line.rstrip()
        if stripped.startswith("#"):
            if current_section:
                sections.append("\n".join(current_section))
                current_section = []
            level = len(stripped.split(" ")[0])
            sections.append(f"\n[제목 {level}단계] {stripped.lstrip('#').strip()}")
        elif "|" in stripped and stripped.strip().startswith("|"):
            current_section.append(f"[표] {stripped}")
        elif stripped:
            current_section.append(f"[내용] {stripped}")

    if current_section:
        sections.append("\n".join(current_section))

    return "\n".join(sections)


def _parse_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    output = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name.lower()
        if "heading" in style:
            level = "".join(filter(str.isdigit, style)) or "1"
            output.append(f"[제목 {level}단계] {text}")
        else:
            output.append(f"[내용] {text}")

    for i, table in enumerate(doc.tables, 1):
        output.append(f"\n[표 {i}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            output.append("  | " + " | ".join(cells) + " |")

    return "\n".join(output)


def _parse_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    output = []

    for i, slide in enumerate(prs.slides, 1):
        output.append(f"\n[슬라이드 {i}]")
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            name = shape.name.lower()
            is_title = "title" in name or shape == slide.shapes.title
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                if is_title:
                    output.append(f"  [제목] {text}")
                else:
                    output.append(f"  [내용] {text}")

        for shape in slide.shapes:
            if shape.shape_type == 19:  # TABLE
                output.append(f"  [표]")
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    output.append("    | " + " | ".join(cells) + " |")

    return "\n".join(output)


def _parse_pdf(file_path: str, pdf_page: int = 0) -> str:
    """pdf_page=0이면 전체, 1 이상이면 해당 페이지만 읽음"""
    import pdfplumber
    output = []

    with pdfplumber.open(file_path) as pdf:
        total_pages = len(pdf.pages)

        if pdf_page > 0:
            if pdf_page > total_pages:
                return f"페이지 범위 초과. 이 PDF는 총 {total_pages}페이지입니다."
            pages_to_read = [(pdf_page, pdf.pages[pdf_page - 1])]
        else:
            pages_to_read = [(i + 1, p) for i, p in enumerate(pdf.pages)]

        for page_num, page in pages_to_read:
            output.append(f"\n[페이지 {page_num}/{total_pages}]")

            tables = page.extract_tables()
            if tables:
                for j, table in enumerate(tables, 1):
                    output.append(f"  [표 {j}]")
                    for row in table:
                        cells = [str(cell or "").strip() for cell in row]
                        output.append("    | " + " | ".join(cells) + " |")

            text = page.extract_text()
            if text:
                for line in text.split("\n"):
                    line = line.strip()
                    if line:
                        output.append(f"  [내용] {line}")

    return "\n".join(output)


@tool
def read_file_structured(file_path: str, pdf_page: int = 0) -> str:
    """파일의 제목, 표, 내용을 구분해서 구조적으로 읽는다.
    지원 형식: pdf, docx, pptx, md
    pdf_page: PDF 전용 — 특정 페이지만 읽을 때 사용 (0=전체, 1 이상=해당 페이지)"""
    if not os.path.exists(file_path):
        return f"파일을 찾을 수 없습니다: {file_path}"

    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".md":
            result = _parse_markdown(file_path)
        elif ext == ".docx":
            result = _parse_docx(file_path)
        elif ext in {".pptx", ".ppt"}:
            result = _parse_pptx(file_path)
        elif ext == ".pdf":
            result = _parse_pdf(file_path, pdf_page=pdf_page)
        else:
            return f"구조 분석을 지원하지 않는 형식입니다: {ext}\n지원 형식: pdf, docx, pptx, md"
    except Exception as e:
        return f"구조 분석 실패: {e}"

    if not result.strip():
        return "구조를 추출할 수 없습니다."

    total = len(result)
    if total > 3000:
        return result[:3000] + f"\n\n...(이하 생략, 전체 {total}자)\n전체 저장은 file_to_note를 사용하세요."
    return result
